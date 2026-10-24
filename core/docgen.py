"""教师助手的文档生成：把 LLM 产出的 JSON 渲染成 .docx / .pptx。

设计说明：
- LLM 只负责「内容」，排版权式由这里用 python-docx / python-pptx 控制，
  这样生成的文件格式稳定，不会因为模型措辞变化而崩掉。
- 中文字体需单独设置 w:eastAsia / a:ea，否则 Word/PPT 里中文会回退成默认字体。
- `parse_json_block` 做了多重容错，因为模型偶尔会用 ```json 包裹或加解释文字。
"""
import json
import re
from pathlib import Path


# ————————————————————————————————————————————————
# JSON 解析容错
# ————————————————————————————————————————————————
def parse_json_block(text: str) -> dict:
    """从模型输出中尽量提取出 JSON 对象。

    依次尝试：直接解析 → 去掉 ```json 包裹 → 截取首尾大括号 → 括号配平扫描。
    全部失败则抛 ValueError（由调用方决定降级方式）。
    """
    raw = (text or "").strip()
    if not raw:
        raise ValueError("模型未返回内容")

    # 1) 直接解析
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        pass

    # 2) 去掉 markdown 代码块标记
    fenced = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.S).strip()
    if fenced != raw:
        try:
            return json.loads(fenced)
        except json.JSONDecodeError:
            pass

    # 3) 截取第一个 { 到最后一个 }
    i, j = fenced.find("{"), fenced.rfind("}")
    if i != -1 and j > i:
        try:
            return json.loads(fenced[i:j + 1])
        except json.JSONDecodeError:
            pass

    # 4) 括号配平扫描（应对尾部有多余文字）
    start = fenced.find("{")
    if start != -1:
        depth, in_str, esc = 0, False, False
        for k in range(start, len(fenced)):
            ch = fenced[k]
            if in_str:
                if esc:
                    esc = False
                elif ch == "\\":
                    esc = True
                elif ch == '"':
                    in_str = False
                continue
            if ch == '"':
                in_str = True
            elif ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    try:
                        return json.loads(fenced[start:k + 1])
                    except json.JSONDecodeError:
                        break

    raise ValueError("模型返回的内容不是有效 JSON，无法生成文件")


def _as_list(v) -> list:
    """把可能是 str / list / None 的字段统一成 list[str]。"""
    if v is None:
        return []
    if isinstance(v, str):
        return [s.strip() for s in re.split(r"[\n;；]+", v) if s.strip()]
    if isinstance(v, (list, tuple)):
        out = []
        for x in v:
            if isinstance(x, dict):
                # 模型有时把列表项写成 {"环节": "...", "内容": "..."}
                out.append("：".join(str(val) for val in x.values() if val))
            else:
                s = str(x).strip()
                if s:
                    out.append(s)
        return out
    return [str(v)]


# ————————————————————————————————————————————————
# DOCX（备课稿 / 教案 / 学情分析共用）
# ————————————————————————————————————————————————
def _docx_set_font(run, name="宋体", size=None, bold=None, color=None):
    """设置中英文字体（中文必须额外写 w:eastAsia）。"""
    from docx.oxml.ns import qn

    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        from docx.oxml import OxmlElement

        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def build_docx(spec: dict, out_path: str, doc_kind: str = "备课稿") -> str:
    """把统一结构的 JSON 渲染成 .docx。返回文件路径。"""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Cm

    doc = Document()

    # 页面与默认样式
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    # 让正文默认也带上东亚字体
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), "宋体")

    # —— 标题 ——
    title = str(spec.get("title") or f"{doc_kind}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    _docx_set_font(p.add_run(title), "黑体", Pt(18), True, RGBColor(0x1F, 0x38, 0x64))

    # 副标题（文档类型）
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(12)
    _docx_set_font(p2.add_run(doc_kind), "楷体", Pt(11), False, RGBColor(0x59, 0x59, 0x59))

    # —— meta 信息栏（2 列表格） ——
    meta = spec.get("meta")
    if isinstance(meta, dict) and meta:
        items = [(str(k), str(v)) for k, v in meta.items() if str(v).strip()]
        if items:
            rows = (len(items) + 1) // 2
            table = doc.add_table(rows=rows, cols=4)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for idx, (k, v) in enumerate(items):
                r, c = idx // 2, (idx % 2) * 2
                cell_k = table.cell(r, c)
                cell_v = table.cell(r, c + 1)
                cell_k.text = ""
                cell_v.text = ""
                _docx_set_font(cell_k.paragraphs[0].add_run(k), "黑体", Pt(9.5), True)
                _docx_set_font(cell_v.paragraphs[0].add_run(v), "宋体", Pt(9.5))
            doc.add_paragraph()

    # —— 正文章节 ——
    sections = spec.get("sections") or []
    if not isinstance(sections, list):
        sections = []
    for n, s in enumerate(sections, 1):
        if not isinstance(s, dict):
            continue
        heading = str(s.get("heading") or f"第{n}部分").strip()
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(10)
        hp.paragraph_format.space_after = Pt(4)
        _docx_set_font(hp.add_run(f"{n}. {heading}"), "黑体", Pt(13), True, RGBColor(0x1F, 0x38, 0x64))

        body = str(s.get("body") or "").strip()
        for para in [t.strip() for t in body.split("\n") if t.strip()]:
            bp = doc.add_paragraph()
            bp.paragraph_format.first_line_indent = Pt(21)  # 首行缩进两字符
            bp.paragraph_format.line_spacing = 1.4
            bp.paragraph_format.space_after = Pt(3)
            _docx_set_font(bp.add_run(para), "宋体", Pt(10.5))

        for it in _as_list(s.get("items")):
            ip = doc.add_paragraph(style="List Bullet")
            ip.paragraph_format.line_spacing = 1.4
            ip.paragraph_format.space_after = Pt(2)
            _docx_set_font(ip.add_run(it), "宋体", Pt(10.5))

    # —— 页脚提示 ——
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _docx_set_font(
        foot.add_run("本文件由校园 AI 助手根据教师上传资料生成，请教师复核后使用。"),
        "楷体", Pt(8.5), False, RGBColor(0x8C, 0x8C, 0x8C),
    )

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


# ————————————————————————————————————————————————
# PPTX（课件）
# ————————————————————————————————————————————————


def _set_ea(run, name="微软雅黑"):
    """给 pptx 的 run 同时设置拉丁与东亚字体，保证中文不回退。"""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    # 移除旧的 ea，避免重复节点
    for old in rPr.findall(qn("a:ea")):
        rPr.remove(old)
    ea = OxmlElement("a:ea")
    ea.set("typeface", name)
    rPr.append(ea)


def _clear_slides(prs):
    """清空模板自带的示例页，只保留母版与版式。

    外部模板通常带几张示例页，若不清空，生成结果末尾会挂着模板作者的样张。
    """
    try:
        sld_id_lst = prs.slides._sldIdLst
        for sld_id in list(sld_id_lst):
            r_id = sld_id.rId
            prs.part.drop_rel(r_id)
            sld_id_lst.remove(sld_id)
    except Exception:  # noqa: BLE001  结构异常时宁可不删，也不让生成失败
        pass


def _pick_blank_layout(prs):
    """挑占位符最少的版式作为空白画布（通常是「空白」版式）。"""
    best, best_n = None, 10 ** 6
    for lay in prs.slide_layouts:
        n = len(lay.placeholders)
        if n < best_n:
            best, best_n = lay, n
    return best or prs.slide_layouts[0]


def _fit_picture(slide, path, box, conv):
    """把图片等比缩放后居中放进 box=(x,y,w,h)，不拉伸变形。"""
    from PIL import Image

    x, y, w, h = box
    IX, IY, IW, IH = conv
    try:
        with Image.open(path) as im:
            iw, ih = im.size
    except Exception:  # noqa: BLE001
        return None
    if not iw or not ih or w <= 0 or h <= 0:
        return None
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    dx = x + (w - dw) / 2
    dy = y + (h - dh) / 2
    try:
        return slide.shapes.add_picture(path, IX(dx), IY(dy), IW(dw), IH(dh))
    except Exception:  # noqa: BLE001
        return None


def build_pptx(spec: dict, out_path: str, template_path: str = None,
               images_map: dict = None) -> str:
    """把课件 JSON 渲染成 16:9 的 .pptx。返回文件路径。

    三种模式：
      · 纯自建（默认）：按内置「学术风」规范从零排版，画布 1280×720 px。
      · 外部模板：传入 template_path，以该 .pptx 为底版 —— 继承它的母版、
        主题配色、字体、背景与装饰，并**按模板画布尺寸等比缩放**所有坐标，
        这样换成 4:3 或别的尺寸模板也不会错位。
      · 自动配图：传入 images_map={页序号: 图片本地路径}，对应页改用
        「左文右图」双栏版式，图片等比缩放居中，不会拉伸变形。

    版式（由模型输出的 type 决定）：
      toc 目录 / section 章节扉页 / content 要点页 / data 数据页
    封面与结束页由程序自动生成，保证一致性。
    """
    from pptx import Presentation
    from pptx.util import Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE

    # —— 画布逻辑尺寸（内部一律按 1280×720 布局，再按模板实际尺寸缩放） ——
    CANVAS_W, CANVAS_H = 1280, 720
    EMU_PER_PX = 9525          # 914400 EMU/in ÷ 96 px/in
    use_template = bool(template_path)

    if use_template:
        prs = Presentation(template_path)
        _clear_slides(prs)
        cur_w = prs.slide_width / EMU_PER_PX
        cur_h = prs.slide_height / EMU_PER_PX
        KX = cur_w / CANVAS_W if cur_w > 0 else 1.0
        KY = cur_h / CANVAS_H if cur_h > 0 else 1.0
    else:
        prs = Presentation()
        KX = KY = 1.0
        prs.slide_width = Emu(int(CANVAS_W * EMU_PER_PX))
        prs.slide_height = Emu(int(CANVAS_H * EMU_PER_PX))

    KF = min(KX, KY)           # 字号按较小比例缩放，防止文字撑破版面

    def IX(px):
        return Emu(int(round(px * KX * EMU_PER_PX)))

    def IY(px):
        return Emu(int(round(px * KY * EMU_PER_PX)))

    def IW(px):
        return IX(px)

    def IH(px):
        return IY(px)

    def F(px):
        """px → pt（96dpi 下 1px = 0.75pt），并按模板缩放。"""
        return Pt(px * 0.75 * KF)

    blank = _pick_blank_layout(prs)

    # —— 学术风配色（用模板时仍沿用，模板主题色体现在背景与母版装饰上） ——
    MAIN = RGBColor(0x1E, 0x4F, 0xA8)
    DARK = RGBColor(0x0E, 0x3F, 0x8C)
    MID = RGBColor(0x3D, 0x7B, 0xD9)
    TXT = RGBColor(0x1A, 0x22, 0x30)
    SUB = RGBColor(0x4A, 0x55, 0x68)
    HINT = RGBColor(0x8B, 0x97, 0xA8)
    LINE = RGBColor(0xD6, 0xDC, 0xE5)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    BG = RGBColor(0xF7, 0xF9, 0xFC)

    PAD_X, PAD_Y = 60, 20
    A_H = 120          # A 标题区高度
    C_Y = 660          # C 页脚区起始

    images_map = images_map or {}

    def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        sp = s.shapes.add_shape(shape, IX(x), IY(y), IW(w), IH(h))
        sp.fill.solid()
        sp.fill.fore_color.rgb = color
        sp.line.fill.background()
        sp.shadow.inherit = False
        return sp

    def text(s, x, y, w, h, content, size, color=TXT, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
        tb = s.shapes.add_textbox(IX(x), IY(y), IW(w), IH(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = str(content)
        _set_ea(r)
        r.font.size = F(size)
        r.font.bold = bold
        r.font.color.rgb = color
        return tb

    def put_image(s, idx, box):
        """若该页有配图则插入并返回 True。"""
        path = images_map.get(idx)
        if not path:
            return False
        return _fit_picture(s, path, box, (IX, IY, IW, IH)) is not None

    # —— 元信息 ——
    course = str(spec.get("course") or "").strip()
    title_text = str(spec.get("title") or "教学课件").strip()
    subtitle = str(spec.get("subtitle") or "").strip()
    audience = str(spec.get("audience") or "").strip()

    raw_slides = spec.get("slides") or []
    if not isinstance(raw_slides, list):
        raw_slides = []
    slides = [x for x in raw_slides if isinstance(x, dict)]
    total_pages = len(slides) + 2  # 封面 + 内容 + 结束页
    page_no = 0

    def footer(s):
        """C 区页脚：左侧课程名，右侧页码。"""
        nonlocal page_no
        rect(s, 0, C_Y, CANVAS_W, 1, LINE)
        left = course or title_text
        text(s, PAD_X, C_Y + 18, 800, 24, left, 14, HINT)
        text(s, CANVAS_W - PAD_X - 200, C_Y + 18, 200, 24,
             f"{page_no:02d} / {total_pages:02d}", 14, MID, align=PP_ALIGN.RIGHT)

    def page_header(s, heading):
        """A 区标题：主标题 + 下方装饰线（长度与标题呼应，不做满宽）。"""
        text(s, PAD_X, 38, CANVAS_W - PAD_X * 2 - 120, 56, heading, 34, MAIN, bold=True)
        rect(s, PAD_X, 100, 56, 4, MID)

    # ==================== 封面 ====================
    s = prs.slides.add_slide(blank)
    page_no += 1
    # 用模板时不再铺满深蓝底色，避免盖住模板自带的封面设计
    if not use_template:
        rect(s, 0, 0, CANVAS_W, CANVAS_H, DARK)
        rect(s, CANVAS_W - 300, 0, 300, CANVAS_H, MAIN)
        rect(s, CANVAS_W - 300, 0, 8, CANVAS_H, MID)

    title_color = WHITE if not use_template else DARK
    sub_color = RGBColor(0xA8, 0xC4, 0xE8) if not use_template else SUB

    if course:
        text(s, PAD_X + 20, 120, 700, 28, course, 20,
             RGBColor(0xA8, 0xC4, 0xE8) if not use_template else MID)
    rect(s, PAD_X + 20, 168, 72, 5, MID)

    tb = s.shapes.add_textbox(IX(PAD_X + 20), IY(210), IW(CANVAS_W - 400), IH(280))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.line_spacing = 1.15
    r = p.add_run()
    r.text = title_text
    _set_ea(r)
    r.font.size = F(72 if len(title_text) <= 14 else 56)
    r.font.bold = True
    r.font.color.rgb = title_color

    if subtitle:
        text(s, PAD_X + 20, 458, CANVAS_W - 400, 40, subtitle, 26, sub_color)
    if audience:
        text(s, PAD_X + 20, 512, CANVAS_W - 400, 30, "适用对象：" + audience, 20,
             RGBColor(0x7F, 0xA0, 0xCC) if not use_template else HINT)

    # ==================== 内容页 ====================
    for idx, item in enumerate(slides):
        s = prs.slides.add_slide(blank)
        page_no += 1
        stype = str(item.get("type") or "content").strip().lower()
        title = str(item.get("title") or "").strip()
        bullets = _as_list(item.get("bullets"))
        note = str(item.get("note") or "").strip()
        has_img = idx in images_map

        # ---------------- 目录页 ----------------
        if stype == "toc":
            rect(s, 0, 0, CANVAS_W, CANVAS_H, WHITE)
            rect(s, 0, 0, 12, CANVAS_H, MAIN)
            page_header(s, title or "目录")
            n = len(bullets)
            if n:
                top = 168
                avail = C_Y - top - 24
                row_h = min(88, max(48, int(avail / max(n, 1))))
                size = 26 if n <= 5 else 22
                for i, b in enumerate(bullets):
                    y = top + i * row_h
                    text(s, PAD_X, y, 76, row_h - 8, f"{i + 1:02d}", size + 8, MID, bold=True)
                    text(s, PAD_X + 84, y + 4, CANVAS_W - PAD_X * 2 - 84, row_h - 8, b, size, TXT)
                    if i < n - 1:
                        rect(s, PAD_X, y + row_h - 10, CANVAS_W - PAD_X * 2, 1, LINE)
            footer(s)

        # ---------------- 章节扉页 ----------------
        elif stype == "section":
            num = str(item.get("number") or "").strip()
            rect(s, 0, 0, CANVAS_W, CANVAS_H, WHITE)
            rect(s, 0, 0, 430, CANVAS_H, DARK)
            if num:
                text(s, 52, 250, 340, 220, num, 160, RGBColor(0x4E, 0x7C, 0xC8), bold=True)
            text(s, 52, 470, 330, 40, "章节", 20, RGBColor(0x7F, 0xA0, 0xCC))

            tx = 500
            tw = CANVAS_W - tx - PAD_X
            text(s, tx, 250, tw, 90, title, 46, MAIN, bold=True)
            rect(s, tx, 352, 72, 5, MID)
            if bullets:
                top = 392
                size = 22 if len(bullets) <= 4 else 20
                for i, b in enumerate(bullets):
                    y = top + i * 42
                    rect(s, tx, y + 11, 7, 7, MID, MSO_SHAPE.OVAL)
                    text(s, tx + 22, y, tw - 22, 32, b, size, SUB)
            footer(s)

        # ---------------- 数据页 ----------------
        elif stype == "data":
            rect(s, 0, 0, CANVAS_W, CANVAS_H, BG)
            page_header(s, title)
            stats = item.get("stats") or []
            if not isinstance(stats, list):
                stats = []
            stats = [x for x in stats if isinstance(x, dict)][:4]
            # 右侧若有配图，KPI 卡片只在左半区排布
            area_w = 620 if has_img else CANVAS_W - PAD_X * 2
            if stats:
                n = len(stats)
                gap = 24
                cols = 2 if (has_img and n > 2) else n
                cw = int((area_w - gap * (cols - 1)) / max(cols, 1))
                cy, ch, row_gap = 160, 170, 16
                for i, st in enumerate(stats):
                    r_i, c_i = divmod(i, cols)
                    cx = PAD_X + c_i * (cw + gap)
                    yy = cy + r_i * (ch + row_gap)
                    rect(s, cx, yy, cw, ch, WHITE)
                    rect(s, cx, yy, cw, 6, MAIN)
                    val = str(st.get("value") or "").strip()
                    lab = str(st.get("label") or "").strip()
                    text(s, cx + 22, yy + 40, cw - 44, 84,
                         val, 60 if len(val) <= 5 else 44, MAIN, bold=True)
                    text(s, cx + 22, yy + 124, cw - 44, 40, lab, 19, SUB)
                rows = (n + cols - 1) // cols
                btop = cy + rows * (ch + row_gap) + 20
            else:
                btop = 176
            put_image(s, idx, (700, 150, 520, 470))
            if bullets:
                size = 24 if len(bullets) <= 4 else 21
                for i, b in enumerate(bullets):
                    y = btop + i * 46
                    if y + 40 > C_Y:
                        break
                    rect(s, PAD_X, y + 13, 8, 8, MID, MSO_SHAPE.OVAL)
                    text(s, PAD_X + 24, y, area_w - 24, 40, b, size, TXT)
            footer(s)

        # ---------------- 要点页（默认） ----------------
        else:
            rect(s, 0, 0, CANVAS_W, CANVAS_H, WHITE)
            rect(s, 0, 0, 12, CANVAS_H, MAIN)
            page_header(s, title)

            # 有配图 → 左文右图双栏；无配图 → 单栏
            text_w = 600 if has_img else CANVAS_W - PAD_X * 2 - 70
            if has_img:
                put_image(s, idx, (700, 152, 520, 468))

            n = max(1, len(bullets))
            top = 168
            avail = C_Y - top - 20
            row_h = min(96, max(52, int(avail / n)))
            size = (26 if n <= 3 else (23 if n <= 5 else 20)) if not has_img else \
                   (23 if n <= 3 else (20 if n <= 5 else 18))
            for i, b in enumerate(bullets):
                y = top + i * row_h
                if y + 48 > C_Y:
                    break
                text(s, PAD_X, y + 2, 62, 40, f"{i + 1:02d}", size + 2, MID, bold=True)
                text(s, PAD_X + 70, y, text_w - 70, row_h - 6, b, size, TXT, spacing=1.35)
                if i < len(bullets) - 1 and y + row_h < C_Y - 10:
                    rect(s, PAD_X + 70, y + row_h - 12, text_w - 70, 1, LINE)
            footer(s)

        if note:
            s.notes_slide.notes_text_frame.text = note

    # ==================== 结束页 ====================
    s = prs.slides.add_slide(blank)
    page_no += 1
    if not use_template:
        rect(s, 0, 0, CANVAS_W, CANVAS_H, DARK)
        rect(s, 0, CANVAS_H - 96, CANVAS_W, 96, MAIN)
    else:
        rect(s, 0, CANVAS_H - 6, CANVAS_W, 6, MAIN)
    text(s, 0, 268, CANVAS_W, 100, "谢谢！", 64,
         WHITE if not use_template else MAIN, bold=True, align=PP_ALIGN.CENTER)
    text(s, 0, 372, CANVAS_W, 40, title_text, 22,
         RGBColor(0xA8, 0xC4, 0xE8) if not use_template else HINT,
         align=PP_ALIGN.CENTER)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path
